#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Livrable 2 — NEXA_Synthese_Marche_Emploi_Marketing_Digital_2026.docx"""
import json, os, sys, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from analyse_mkt import FAMILLE_INDIC, HYPOTHESES, SCENARIOS, CRITERES, OPTIONS, scores
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
BASE = os.path.join(ROOT, "collecte_mkt")
D = json.load(open(os.path.join(BASE, "mkt_consolide.json"), encoding="utf-8"))
ETUDES = [json.loads(l) for l in open(os.path.join(BASE, "mkt_etudes.jsonl"), encoding="utf-8") if l.strip()]
NAT = D["national"]; META = D["meta"]
NAVY = RGBColor(0x1F, 0x38, 0x64); BLUE = RGBColor(0x2E, 0x75, 0xB6); GREY = RGBColor(0x60, 0x60, 0x60)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)

def shade(cell, hexcolor):
    el = OxmlElement("w:shd"); el.set(qn("w:fill"), hexcolor); cell._tc.get_or_add_tcPr().append(el)

def H1(t):
    p = doc.add_paragraph(); p.space_before = Pt(16); p.space_after = Pt(6)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = NAVY
    return p
def H2(t):
    p = doc.add_paragraph(); p.space_before = Pt(11); p.space_after = Pt(4)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = BLUE
    return p
def H3(t):
    p = doc.add_paragraph(); p.space_before = Pt(8); p.space_after = Pt(3)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(10.5)
    return p
def P(t, size=10.5, bold=False, italic=False, color=None, align=None, after=4):
    p = doc.add_paragraph(); p.space_after = Pt(after)
    if align: p.alignment = align
    r = p.add_run(t); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p
def BUL(t, size=10):
    p = doc.add_paragraph(style="List Bullet"); p.space_after = Pt(2)
    r = p.add_run(t); r.font.size = Pt(size); return p
def link(paragraph, url, text=None):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "17"); rPr.append(sz)
    r.append(rPr); t = OxmlElement("w:t"); t.text = text or url; t.set(qn("xml:space"), "preserve")
    r.append(t); h.append(r); paragraph._p.append(h)
def TAB(headers, rows, widths=None, fontsize=8.5, headcolor="1F3864"):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(fontsize)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, headcolor)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(fontsize)
    if widths:
        for r_ in t.rows:
            for i, w in enumerate(widths):
                r_.cells[i].width = Cm(w)
    return t

# =============================================================== COUVERTURE
P("NEXA DIGITAL SCHOOL", 12, True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_after = Pt(4)
r = p.add_run("OBSERVATOIRE DU MARCHÉ DE L'EMPLOI\nDU MARKETING DIGITAL EN FRANCE")
r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
P("Quelle orientation NEXA doit-elle donner à sa filière Marketing Digital ?", 13, True,
  color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
P("Synthèse pour la Direction Générale, la Direction Marketing et la Direction Pédagogique", 10.5,
  italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("Collecte et analyse du 7 septembre 2026", 10, italic=True, color=GREY,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=16)

# ---- encadré conclusion
t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
c = t.rows[0].cells[0]; shade(c, "1F3864"); c.text = ""
pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = pp.add_run("CONCLUSION"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
pp2 = c.add_paragraph(); pp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = pp2.add_run("NEXA doit ADOPTER UN MODÈLE HYBRIDE SPÉCIALISÉ : conserver en Bachelor un socle "
                "généraliste mais réellement technicisé — c'est lui qui capte l'alternance et l'accès junior — "
                "et transformer le Mastère en deux spécialisations techniques, Growth & Performance d'une part, "
                "CRM, Data & Marketing Automation d'autre part, en supprimant le parcours Social Media / "
                "Community Management autonome et en traitant l'IA comme une couche de compétences "
                "obligatoire et transversale, jamais comme un parcours dédié.")
r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
doc.add_paragraph()

P(f"Base de l'étude : {META['volume_brut']} offres collectées sur 4 plateformes, "
  f"{META['offres_uniques']} offres uniques après dédoublonnage, dont {META['offres_dans_perimetre']} dans le périmètre "
  f"marketing digital ({META['offres_coeur_marketing_digital']} hors métiers adjacents) ; "
  f"{META['nb_volumes']} stocks d'offres datés relevés sur les jobboards (2024-2026) ; "
  f"{META['nb_etudes']} sources documentaires (Apec, France Travail, Fevad, SRI-UDECAM-Oliver Wyman, "
  f"Alliance Digitale/EY, Dares, observatoires salariaux).", 9.5, italic=True, color=GREY)
P("Sources par plateforme : " + " · ".join(f"{k} {v}" for k, v in sorted(META["sources"].items(), key=lambda x:-x[1])),
  9.5, italic=True, color=GREY)
doc.add_page_break()

# =============================================================== 1. MARCHÉ
H1("1. État du marché national")
P("Le marché français du marketing digital ne recule pas : il croît, mais il se polarise.", 11, True)
P("La filière marketing digital française représente 14,4 milliards d'euros de revenus directs et 5,6 milliards "
  "de valeur ajoutée en 2024, soit davantage que la presse ou la musique, pour environ 310 000 emplois répartis sur "
  "plus de 18 000 entreprises. Sa croissance est cinq fois supérieure à celle du PIB français et elle consacre "
  "1,6 milliard d'euros à la R&D, soit 11,5 % de son chiffre d'affaires — un niveau inhabituellement élevé pour un "
  "secteur de services (Alliance Digitale / EY, 2025). Près de la moitié de ces emplois sont situés hors "
  "Île-de-France.")
P("Les marchés qui financent ces emplois progressent également. Le marché de la publicité digitale a crû de 11 % en "
  "2025 et de 12 % au premier semestre 2026, pour atteindre 6,689 milliards d'euros sur le seul S1 2026 ; le Social "
  "(33 % du marché, +16 %) se rapproche du Search (41 %) et le retail media progresse de 18 % "
  "(SRI / UDECAM / Oliver Wyman, 36e Observatoire de l'e-pub). L'e-commerce français atteint 196,4 milliards d'euros "
  "en 2025 (+7 %, en ralentissement par rapport aux +9,6 % de 2024) et représente 234 000 emplois, en hausse de 9 % "
  "sur un an (Fevad, Chiffres clés 2026).")
P("Le contexte de l'emploi cadre est en amélioration relative : l'Apec prévoit 305 800 recrutements de cadres en 2026, "
  "soit +4 % par rapport à 2025, les fonctions informatique, études-R&D et commercial-marketing concentrant à elles "
  "seules plus de la moitié des recrutements attendus. Ce signal doit toutefois être lu avec prudence : l'enquête "
  "Besoins en Main-d'Œuvre 2026 de France Travail recense près de 2,3 millions d'intentions de recrutement tous "
  "métiers confondus, en baisse de 6,5 % par rapport à 2025. Le marketing digital évolue donc mieux que la moyenne "
  "du marché du travail, mais dans un environnement macro-économique dégradé.")

H2("Ce que montre l'échantillon d'offres")
TAB(["Indicateur", "Valeur observée", "Lecture"],
 [["Offres uniques dans le périmètre", f"{META['offres_dans_perimetre']}", "Base de tous les pourcentages qui suivent"],
  ["Taux de doublons multi-plateformes", f"{META['taux_doublons']} %", "Faible : les plateformes se recoupent peu sur les offres individuelles indexées"],
  ["Part CDI", f"{NAT['contrats_pct'].get('CDI',0)} % ({NAT['contrats'].get('CDI',0)} offres)", "Contrat dominant lorsqu'il est renseigné"],
  ["Part alternance", f"{NAT['contrats_pct'].get('ALTERNANCE',0)} % ({NAT['contrats'].get('ALTERNANCE',0)} offres)", "Deuxième contrat du marché — voie d'entrée dominante"],
  ["Part stage", f"{NAT['contrats_pct'].get('STAGE',0)} %", "Marginale par rapport à l'alternance"],
  ["Contrat non renseigné", f"{NAT['contrats_pct'].get('NC',0)} %", "Limite de méthode : les extraits indexés ne le précisent pas toujours"],
  ["Débutants + juniors", f"{NAT['debutant_junior_pct']} %", "Accès junior réel mais très inégal selon les familles"],
  ["Séniorité non renseignée", f"{NAT['seniorite_pct'].get('NC',0)} %", "Limite de méthode majeure — à ne pas surinterpréter"],
  ["Technicité moyenne (échelle 1-4)", f"{NAT['technicite_moy']}", "Le marché est majoritairement de niveau 2 (expertise canal)"],
  ["Offres mentionnant explicitement l'IA", f"{NAT['ia_pct']} %", "Converge avec la mesure Apec (2 % des offres commercial-marketing)"],
  ["Salaire annuel brut médian affiché", f"{NAT['salaire_median']} € (n={NAT['salaire_n']})", "Effectif trop faible pour conclure — voir la section Salaires"]],
 [5.2, 4.6, 7.2])
P("Limite de méthode à garder en tête pour toute la suite : les données proviennent des titres, URL et extraits de "
  "pages publiques indexées par un moteur de recherche, aucun accès direct aux sites ni aux API n'ayant été possible. "
  "Les champs absents des extraits sont notés NC et n'ont jamais été déduits. Les comptes de compétences sont donc "
  "des bornes basses, et les parts de contrat et de séniorité sont calculées sur l'ensemble des offres, y compris "
  "celles où l'information est absente.", 9, italic=True, color=GREY)

# =============================================================== 2. ÉVOLUTION
H1("2. Évolution de la demande : ce qui est démontré et ce qui ne l'est pas")
P("Point de rigueur méthodologique. Les séries historiques comparables métier par métier n'ont pas pu être "
  "reconstituées : les pages de listes d'offres sont indexées à des dates hétérogènes et sous des libellés de "
  "requête différents, ce qui interdit toute comparaison année contre année à périmètre constant. L'étude ne peut "
  "donc PAS affirmer qu'un métier donné a progressé ou reculé de X % en un an. Elle établit en revanche une "
  "structure de marché à septembre 2026, et croise cette structure avec les évolutions publiées par les organismes "
  "de référence.", 10, italic=True)
H2("Ce qui est démontré")
BUL("Les marchés amont progressent : publicité digitale +11 % en 2025 et +12 % au S1 2026 ; e-commerce +7 % de chiffre d'affaires et +9 % d'emplois en 2025.")
BUL("La filière croît cinq fois plus vite que le PIB et pèse 310 000 emplois (Alliance Digitale / EY).")
BUL("Les intentions de recrutement de cadres repartent : +4 % en 2026 selon l'Apec, le commercial-marketing figurant parmi les trois premières fonctions.")
BUL("Les offres de la fonction commerciale mentionnant l'intelligence artificielle ont progressé de 62 % entre 2022 et 2025 (Apec) — une dynamique forte, à partir d'un niveau très bas.")
BUL("Les salaires des métiers techniques progressent : SEO manager confirmé à 47 855 € de salaire médian, en hausse de 4,2 % sur un an ; spécialiste SEO à 48 822 €, +3,2 % (Observatoire Seobooster 2026).")
H2("Ce qui n'est pas démontré")
BUL("Le rythme de progression comparé des métiers généralistes et des métiers spécialisés (hypothèse H1) : la structure du marché est établie, la vitesse ne l'est pas.")
BUL("La substitution des CDI juniors par l'alternance (hypothèse H5) : seule l'abondance de l'alternance est démontrée.")
BUL("Toute évolution régionale ou locale année contre année.")

# =============================================================== 3. MÉTIERS
H1("3. Métiers qui progressent, métiers qui reculent")
H2("Structure observée du marché par famille de métiers")
rows = []
for f, b in sorted(D["par_famille"].items(), key=lambda x: -x[1]["n"]):
    ind = FAMILLE_INDIC.get(f, ("NC","NC","NC","NC",""))
    rows.append([f, b["n"], f"{round(100*b['n']/META['offres_dans_perimetre'],1)} %",
                 b["technicite_moy"], f"{b['alternance_pct']} %", f"{b['debutant_junior_pct']} %",
                 ind[0], ind[1], ind[2]])
TAB(["Famille de métiers","Offres","Part","Technicité","Alternance","Déb.+juniors","Tension","Qualité du débouché","Exposition IA"],
    rows, [3.6,1.2,1.1,1.4,1.5,1.6,2.2,2.4,1.9], 8)
P("Lecture : la technicité est notée de 1 (exécution) à 4 (MarTech et systèmes). La corrélation est frappante et "
  "constante — plus une famille est technique, moins elle recrute de juniors et d'alternants, mais plus elle est "
  "résistante à l'automatisation et mieux elle rémunère.", 9, italic=True, color=GREY)

H2("Les métiers qui progressent")
BUL("Growth, acquisition et performance. Première famille spécialisée avec 85 offres uniques (17,2 % du périmètre). Les stocks jobboards confirment : 1 842 offres Growth Marketing en France sur Indeed (août 2026) et 1 295 en Île-de-France (septembre 2026). C'est la famille directement financée par la croissance du marché publicitaire digital.")
BUL("CRM, lifecycle et marketing automation. 36 offres uniques, mais la compétence CRM est la deuxième la plus citée de tout le périmètre (11,5 % des offres) et les stocks sont élevés : plus de 2 000 offres CRM Manager en France et 800 à Paris (septembre 2026), 759 offres Marketing Automation (mai 2026). Technicité 2,50, faible exposition à l'automatisation.")
BUL("Data marketing, analytics et tracking. 31 offres, technicité 2,45, avec des exigences explicites en SQL, BigQuery, GA4, GTM et Looker Studio. Des intitulés hybrides apparaissent : Senior Data Analyst SEO chez Havas, Lead Marketing Data Analyst chez Doctolib, Marketing Data Analyst chez Groupe Positive.")
BUL("E-commerce. 28 offres, portées par un secteur à 196,4 milliards d'euros et 234 000 emplois en croissance de 9 %.")
BUL("GEO, AEO et AI Search — émergent. Sept offres citent explicitement ces compétences et des intitulés réels existent : Consultant SEO/GEO chez datashake et Ad's up, SEO SEA GEO AI Visibility Manager E-commerce chez Quincaillerie Angles, Community Manager spécialiste Agents IA & GEO/SEO. Le signal est réel mais le volume reste très faible.")

H2("Les métiers qui reculent ou se fragilisent")
BUL("Social media et community management. 39 offres pour la technicité la plus faible du périmètre (1,03). France Travail ne recense que 317 offres de community manager en 2026 face à un vivier de candidats très large ; la saturation est documentée par les observateurs du métier. Par ailleurs, 52 % des annonceurs prévoient de réduire leur budget média en 2026. C'est le seul segment de l'étude qui réunit tous les critères de fragilité.")
BUL("Production de contenus génériques. 76 % des marketeurs déclarent utiliser l'IA générative pour la création de contenu. Les postes de contenu se maintiennent sur la stratégie éditoriale, le brand content et l'influence — pas sur la production.")
BUL("Marketing digital généraliste d'exécution. 136 offres de l'échantillon (27,5 %) relèvent du niveau de technicité 1 et décrivent précisément les tâches que les études identifient comme les plus automatisables : publication, mise à jour de contenus, emailing simple, reporting manuel, coordination. Les salaires observés sur ces postes sont les plus bas du périmètre (20 000 à 30 000 €).")

# =============================================================== 4. GEN vs SPE
H1("4. Généralistes contre spécialistes : le résultat central de l'étude")
spec = ["ACQUISITION_PERFORMANCE","SEO_SEA","CRM_LIFECYCLE","DATA_ANALYTICS_CRO","ECOMMERCE","MARTECH_MARKETING_OPS","PRODUCT_MARKETING"]
n_spec = sum(D["par_famille"].get(x, {}).get("n", 0) for x in spec)
n_gen = D["par_famille"].get("COEUR_GENERALISTE", {}).get("n", 0)
P(f"Les familles spécialisées totalisent {n_spec} offres uniques, soit {round(100*n_spec/META['offres_dans_perimetre'],1)} % du périmètre, "
  f"contre {n_gen} offres pour le cœur généraliste ({round(100*n_gen/META['offres_dans_perimetre'],1)} %). "
  "Le marché français du marketing digital est donc majoritairement spécialisé.", 11, True)
P("Mais ce constat, pris seul, conduirait NEXA à une erreur stratégique. Car la répartition des débouchés juniors "
  "est exactement inverse.")
TAB(["Famille","Offres","Technicité","Part d'alternance","Part débutants + juniors"],
 [[f, b["n"], b["technicite_moy"], f"{b['alternance_pct']} %", f"{b['debutant_junior_pct']} %"]
  for f, b in sorted(D["par_famille"].items(), key=lambda x: -x[1]["debutant_junior_pct"])],
 [5.0, 1.6, 2.0, 3.4, 4.5], 8.5)
P("Le cœur généraliste est à la fois la famille la plus volumineuse (111 offres), la première pourvoyeuse "
  "d'alternance (40,5 %) et la plus accessible aux débutants (47,7 %) — et en même temps la moins technique (1,42) "
  "et la plus exposée à l'automatisation. À l'inverse, les familles les plus techniques et les mieux rémunérées "
  "n'offrent quasiment aucune porte d'entrée : 0 % d'alternance en MarTech, en Product Marketing et sur les métiers "
  "IA émergents, 3,2 % en data marketing.", 10.5, True)
P("C'est la tension centrale que NEXA doit arbitrer : le marché où se trouve la valeur n'est pas le marché où se "
  "trouvent les débouchés d'entrée. Un Bachelor purement technique n'aurait ni alternance ni placement ; un Mastère "
  "généraliste n'aurait ni valeur ni différenciation. Aucun positionnement unique ne répond aux deux contraintes — "
  "seule une architecture à deux étages le permet.", 11, True, color=NAVY)

# =============================================================== 5. RÉGIONS
H1("5. Répartition régionale")
rows = [[r, b["n"], f"{b['part_nationale']} %", b["technicite_moy"], f"{b['alternance_pct']} %",
         " · ".join(f"{m} ({n})" for m, n in b["metiers_top"][:3])]
        for r, b in sorted(D["par_region"].items(), key=lambda x: -x[1]["n"])]
TAB(["Région","Offres","Part nationale","Technicité","Alternance","Métiers dominants"], rows,
    [4.0,1.2,1.8,1.5,1.5,7.5], 8)
P("L'Île-de-France concentre 52,9 % des offres du périmètre, devant l'Auvergne-Rhône-Alpes. Cette concentration "
  "est encore plus marquée sur les métiers rares : RevOps, Customer Data Platform, Product Marketing et Paid Media "
  "en agence sont presque exclusivement franciliens. En régions dominent le marketing digital généraliste, "
  "l'e-commerce, le SEO/SEA et le CRM.")
P("Nuance importante : l'étude Alliance Digitale / EY établit que près de la moitié des 310 000 emplois de la filière "
  "sont situés hors Île-de-France. La surreprésentation francilienne observée sur les jobboards reflète donc "
  "surtout la concentration des postes qualifiés, des agences et des sièges — pas la totalité de l'emploi.")

# =============================================================== 6. VILLES
H1("6. Les six campus NEXA")
STOCKS = {
 "Paris / Île-de-France":"Paris 2 000+ (14/07/2026) · IDF 3 000+ (18/04/2026) · alternance IDF 1 254 (26/06/2026)",
 "Lyon métropole":"Lyon 317 (02/05/2026) · 258 (28/05/2026)",
 "Lille métropole":"Lille 113 (16/01/2026)",
 "Bordeaux métropole":"Bordeaux 90 (19/07/2026)",
 "Nantes métropole":"Nantes 82 (21/04/2026)",
 "Marseille - Aix":"Marseille 92 (03/09/2026)",
 "National / distanciel":"À distance 100+ (19/05/2026) · full remote 75+ (11/08/2026)",
}
RECO = {
 "Paris / Île-de-France":"Offre complète : Bachelor socle technicisé + les deux Mastères + modules MarTech/RevOps et Product Marketing.",
 "Lyon métropole":"Bachelor socle technicisé + Mastère Growth & Performance. Écosystème d'agences dense.",
 "Lille métropole":"Bachelor socle technicisé, coloration CRM et e-commerce (retail, distribution). Mastère en distanciel.",
 "Bordeaux métropole":"Bachelor socle technicisé, coloration e-commerce et acquisition. Mastère en distanciel.",
 "Nantes métropole":"Marché le plus étroit des six campus. Bachelor socle technicisé uniquement, coloration e-commerce et SEO/SEA. Ne pas ouvrir de Mastère spécialisé en présentiel.",
 "Marseille - Aix":"Bachelor socle technicisé, coloration e-commerce et social ads. Marché d'alternance actif. Mastère en distanciel.",
 "National / distanciel":"Levier décisif pour rentabiliser les Mastères spécialisés à faible volume local. Des postes techniques en full remote existent réellement.",
}
rows = []
for v, b in D["par_ville"].items():
    rows.append([v, STOCKS.get(v,"NC"), b.get("n",0), f"{b.get('alternance_pct',0)} %",
                 b.get("technicite_moy","NC"), RECO.get(v,"NC")])
TAB(["Campus","Stock d'offres jobboard (daté)","Offres échantillon","Alternance","Technicité","Recommandation"],
    rows, [2.8,4.6,1.6,1.4,1.4,5.7], 8)
P("La hiérarchie des marchés locaux est nette : Paris, puis Lyon très nettement en second, puis un groupe resserré "
  "Lille / Marseille / Bordeaux autour de 90 à 113 offres, et enfin Nantes, le marché le plus étroit avec 82 offres. "
  "Cette hiérarchie commande directement la carte des formations : les spécialisations de Mastère les plus techniques "
  "ne sont soutenables en présentiel qu'à Paris et à Lyon.")

# =============================================================== 7. CONTRATS
H1("7. Contrats, alternance, accès junior et séniorité")
H2("Répartition des contrats")
TAB(["Type de contrat","Offres","Part du périmètre"],
 [[k, v, f"{NAT['contrats_pct'][k]} %"] for k, v in sorted(NAT["contrats"].items(), key=lambda x:-x[1])],
 [5.0,3.0,4.0], 9)
H2("L'alternance, actif le plus solide de la filière")
P("L'alternance représente 20,8 % des offres du périmètre, soit 103 offres, et c'est de loin le canal d'entrée "
  "dominant. Les volumes de marché sont considérables et convergents : 1 177 offres d'alternance en marketing "
  "digital sur HelloWork (août 2026), plus de 1 000 sur Indeed (septembre 2026), 1 254 pour la seule Île-de-France "
  "(juin 2026), auxquelles s'ajoutent 459 offres d'alternance chargé de marketing, 375 en community manager, "
  "plus de 100 en responsable marketing digital et 45 en SEO manager.")
P("Mais cette alternance est très inégalement répartie. Elle est concentrée sur le généraliste (40,5 % des offres "
  "de la famille), le marketing adjacent (34,7 %) et le social media (25,6 %). Elle est quasi inexistante sur les "
  "familles les plus techniques : 3,2 % en data marketing, 0 % en MarTech, en Product Marketing et sur les métiers "
  "IA émergents.")
P("Conséquence directe pour NEXA : le socle généraliste n'est pas un héritage à assumer, c'est la condition d'accès "
  "au vivier d'alternance. Le supprimer reviendrait à supprimer le modèle économique du Bachelor.", 10.5, True)
H2("Accès junior et séniorité")
P("25,5 % des offres du périmètre sont explicitement de niveau débutant et 4,2 % de niveau junior, mais 52,9 % ne "
  "précisent pas le niveau d'expérience dans les extraits indexés — cette dernière proportion invite à ne pas "
  "surinterpréter les valeurs absolues. Les écarts relatifs entre familles restent, eux, parfaitement lisibles et "
  "sont présentés au chapitre 4.")
P("Les exigences observées sur les métiers techniques sont sans ambiguïté : 5 ans minimum pour un Web Analyst, "
  "6 ans pour un Marketing Automation Manager, 7 ans pour un CRM Lead, 8 ans pour un Head of Growth, 10 à 15 ans "
  "pour un Head of Marketing. Le Product Marketing exige systématiquement 3 ans et plus et n'affiche aucune offre "
  "junior ni alternance dans l'échantillon.")

# =============================================================== 8. SALAIRES
H1("8. Salaires")
P("Seules 13 offres de l'échantillon affichent une rémunération annuelle exploitable — un effectif beaucoup trop "
  "faible pour produire des médianes par métier. Les valeurs issues des offres sont donc présentées à titre "
  "indicatif, et l'analyse s'appuie principalement sur les études salariales publiées.", 10, italic=True)
TAB(["Statut de la donnée","Métier","Fourchette","Source"],
 [["SALAIRE_AFFICHE_DANS_OFFRE","Ensemble du périmètre", f"médiane {NAT['salaire_median']} €, Q1 {NAT['salaire_q1']} €, Q3 {NAT['salaire_q3']} € (n=13)","Échantillon d'offres 2026"],
  ["SALAIRE_AFFICHE_DANS_OFFRE","Communication digitale / webmarketing (province)","20 000 à 30 000 €","Offres Sèvremoine, Foix, Cournon-d'Auvergne, Périgueux"],
  ["SALAIRE_AFFICHE_DANS_OFFRE","Traffic Manager / Social Ads","28 000 à 48 000 €","Offres Valence, Qeads (full remote), Nantes (34-38 k€)"],
  ["SALAIRE_AFFICHE_DANS_OFFRE","Trade marketing (Bac+5)","35 000 à 42 000 €","Offre ZWILLING STAUB France"],
  ["SALAIRE_AFFICHE_DANS_OFFRE","Growth Marketing Manager (full remote)","55 000 à 65 000 €","Offre Indeed, télétravail total France"],
  ["ESTIMATION_DE_CABINET","Responsable marketing digital","40 000 à 60 000 €","Robert Half 2026"],
  ["SALAIRE_ISSU_D_ETUDE","Growth Marketing Manager","40-50 k€ débutant · 50-65 k€ confirmé · 80-90 k€ senior","Licorne Society 2026"],
  ["SALAIRE_ISSU_D_ETUDE","Traffic Manager","35-42 k€ débutant · 42-50 k€ à 3-5 ans · plus de 60 k€ senior","Licorne Society 2026"],
  ["SALAIRE_ISSU_D_ETUDE","SEO manager confirmé","47 855 € médian, en hausse de 4,2 % sur un an","Observatoire Seobooster 2026"],
  ["SALAIRE_ISSU_D_ETUDE","Spécialiste SEO","48 822 € moyen, en hausse de 3,2 % sur un an","Observatoire Seobooster 2026"],
  ["SALAIRE_ISSU_D_ETUDE","SEO confirmé Paris CDI","65 000 € médian (58-78 k€), jusqu'à 85-95 k€ avec primes anglais et expertise IA/GEO","Observatoire Seobooster 2026"]],
 [3.6,3.4,5.2,3.0], 8)
P("Le message est cohérent avec le reste de l'étude : l'écart de rémunération entre l'exécution généraliste "
  "(20 000 à 30 000 €) et l'expertise technique senior (60 000 à 90 000 €) est d'un facteur trois. La prime "
  "explicite à l'expertise IA et GEO relevée sur les profils SEO parisiens est le premier signal salarial mesurable "
  "de la valorisation de ces compétences.")

# =============================================================== 9. COMPÉTENCES
H1("9. Compétences, outils et niveau de technicité")
H2("Les compétences les plus demandées")
top = sorted(D["competences"].items(), key=lambda x: -x[1]["n"])[:18]
TAB(["Compétence","Offres","Part","Famille"],
 [[n, v["n"], f"{v['pct']} %", v["famille"]] for n, v in top], [6.6,1.5,1.5,4.5], 8.5)
P("Ces comptes sont des bornes basses : ils ne retiennent que les compétences explicitement citées dans le titre ou "
  "l'extrait indexé de l'offre.", 9, italic=True, color=GREY)
H2("Répartition par niveau de technicité")
nt = NAT["technicite"]; tot = sum(nt.values())
TAB(["Niveau","Définition","Offres","Part"],
 [["NIVEAU 1","Généraliste / exécution : publication, rédaction simple, emailing simple, coordination, mise à jour de contenus", nt.get("1",nt.get(1,0)), f"{round(100*nt.get('1',nt.get(1,0))/tot,1)} %"],
  ["NIVEAU 2","Expertise canal : SEO, SEA, Paid Social, CRM, e-commerce, content strategy, influence", nt.get("2",nt.get(2,0)), f"{round(100*nt.get('2',nt.get(2,0))/tot,1)} %"],
  ["NIVEAU 3","Performance / data / automation : GA4, GTM, tracking, attribution, CRO, CRM automation, dashboards, SQL", nt.get("3",nt.get(3,0)), f"{round(100*nt.get('3',nt.get(3,0))/tot,1)} %"],
  ["NIVEAU 4","MarTech / data / systèmes : architecture CRM, CDP, server-side, data warehouse, API, Marketing Ops, RevOps", nt.get("4",nt.get(4,0)), f"{round(100*nt.get('4',nt.get(4,0))/tot,1)} %"]],
 [1.8,9.6,1.5,1.5], 8.5)
P("Le marché français du marketing digital est aujourd'hui massivement de niveau 2 (62,2 % des offres) avec une "
  "base de niveau 1 encore importante (27,5 %). Les niveaux 3 et 4 ne représentent que 10,3 % des offres — mais ce "
  "sont eux qui portent les salaires, la résilience à l'automatisation et l'exigence de Bac+5. Le niveau 4 en "
  "particulier (11 offres) constitue un marché étroit, réservé à des profils de 3 à 7 ans d'expérience : il ne peut "
  "pas fonder un parcours de formation autonome.")

# =============================================================== 10. FAMILLES DÉTAIL
LIBELLES_FAMILLE = {
 "COEUR_GENERALISTE": "Cœur généraliste du marketing digital",
 "ACQUISITION_PERFORMANCE": "Acquisition, performance et growth",
 "SEO_SEA": "SEO, SEA et search",
 "SOCIAL_MEDIA": "Social media et community management",
 "CONTENT_BRAND": "Content, brand et influence",
 "CRM_LIFECYCLE": "CRM, lifecycle et fidélisation",
 "ECOMMERCE": "E-commerce",
 "DATA_ANALYTICS_CRO": "Data, analytics, tracking et CRO",
 "MARTECH_MARKETING_OPS": "MarTech, Marketing Ops et RevOps",
 "PRODUCT_MARKETING": "Product Marketing",
 "METIER_IA_EMERGENT": "Métiers émergents liés à l'IA",
 "METIER_ADJACENT": "Métiers adjacents",
}
H1("10. Lecture détaillée par domaine")
for f in ["ACQUISITION_PERFORMANCE","SEO_SEA","CRM_LIFECYCLE","ECOMMERCE","DATA_ANALYTICS_CRO",
          "MARTECH_MARKETING_OPS","PRODUCT_MARKETING","SOCIAL_MEDIA","CONTENT_BRAND","METIER_IA_EMERGENT"]:
    b = D["par_famille"].get(f, {})
    ind = FAMILLE_INDIC.get(f, ("NC","NC","NC","NC","NC"))
    H2(LIBELLES_FAMILLE.get(f, f) + f" — {b.get('n',0)} offres uniques")
    P(ind[4], 10)
    P(f"Tension : {ind[0]} · Qualité du débouché : {ind[1]} · Exposition à l'automatisation : {ind[2]} · "
      f"Pérennité à 3-5 ans : {ind[3]}", 9, italic=True, color=GREY)

# =============================================================== 11. IA
H1("11. La transformation par l'intelligence artificielle")
H2("L'IA transforme les tâches, elle ne supprime pas les postes")
P("Les usages sont massifs et documentés : 75 % des marketeurs déclarent utiliser l'IA générative dans leur travail "
  "quotidien et 76 % pour la création de contenu ; 88 % des équipes marketing déclarent l'utiliser au quotidien ; "
  "82 % des entreprises d'e-commerce l'utilisaient déjà en 2025 (Fevad). Mais la supervision humaine reste la règle : "
  "86 % des marketeurs retravaillent systématiquement les contenus générés avant publication et 92 % imposent une "
  "validation humaine avant toute diffusion sur les réseaux sociaux.")
P("L'Apec établit le mécanisme : l'IA automatise de nombreuses tâches et redéploie la valeur ajoutée des cadres du "
  "commercial-marketing vers l'analyse, la personnalisation et la stratégie de la relation client. Elle identifie "
  "deux compétences émergentes précises : la maîtrise des outils génératifs et la capacité à les intégrer dans des "
  "processus métier.")
H2("Tâches exposées et compétences dont la valeur augmente")
TAB(["Tâches les plus exposées à l'automatisation","Compétences dont la valeur augmente"],
 [["Rédaction de contenus génériques","Analyse et interprétation des données"],
  ["Déclinaisons publicitaires et production créative","Personnalisation et stratégie de la relation client"],
  ["Publication et animation simple des réseaux sociaux","Qualité de la donnée, tracking et attribution"],
  ["Reporting manuel et synthèse de performances","Architecture CRM et orchestration de systèmes"],
  ["Recherche de mots-clés basique et briefs simples","Expérimentation et CRO"],
  ["Création d'emails simples et segmentation élémentaire","Contrôle des sorties IA, gouvernance, RGPD"],
  ["Production de présentations marketing","Pilotage budgétaire et compréhension business"]],
 [8.0,8.0], 9)
H2("Le « marketeur augmenté par l'IA » a-t-il une réalité observable ?")
P("Oui, mais très marginale — et c'est le résultat le plus important de cette étude pour NEXA.", 11, True)
P(f"Seules {NAT['ia_pct']} % des offres du périmètre mentionnent explicitement l'IA, soit 12 offres sur 495. "
  "Ce chiffre converge remarquablement avec une mesure totalement indépendante : l'Apec établit que les offres "
  "mentionnant l'IA ne représentent que 2 % des offres de la fonction commercial-marketing, alors même que cette "
  "fonction concentre 13 % de l'ensemble des offres cadres mentionnant l'IA, en deuxième position derrière "
  "l'informatique. Deux méthodes différentes aboutissent au même ordre de grandeur.")
P("Des intitulés réels existent bel et bien — dix offres, soit 2,0 % du périmètre : Consultant SEO/GEO, SEO SEA GEO "
  "AI Visibility Manager E-commerce, AI Content Manager, Community Manager spécialiste Agents IA & GEO/SEO, Growth "
  "Manager IA, RevOps Manager AI Data & Automation. La dynamique est réelle : les offres de la fonction commerciale "
  "mentionnant l'IA ont progressé de 62 % entre 2022 et 2025. Mais le niveau de départ est très bas.")
P("Une contradiction doit être signalée explicitement. La recherche ciblée sur les offres mentionnant les agents IA, "
  "le no-code, Make, n8n et Zapier fait remonter huit offres, dont six sont des postes techniques ou produit — "
  "Product Builder IA, Ingénieur IA & Automatisation, Agent Builder GenAI, Développeur intégrations API et agents "
  "MCP — et non des postes marketing. Le marché de l'orchestration d'agents IA et de l'automatisation no-code existe "
  "bien en France en 2026, mais il est aujourd'hui capté par des profils techniques, pas par des marketeurs.", 10.5, True)
P("Conséquence pour NEXA : créer un parcours intitulé « Marketeur augmenté par l'IA » ne serait pas soutenu par la "
  "demande employeur observable. L'IA doit être une couche de compétences transversale et obligatoire — outils "
  "génératifs, contrôle des sorties, automatisation no-code, GEO/AEO — intégrée à chaque spécialisation, et non un "
  "parcours affiché. Une réserve doit être formulée : les extraits indexés ne restituent qu'une partie du contenu "
  "des offres, et 2,4 % est donc une borne basse.", 10.5, True, color=NAVY)

# =============================================================== 12. HYPOTHÈSES
doc.add_page_break()
H1("12. Test des hypothèses")
for code, enonce, res, preuves, chiffres, sources, contra, conf, implic in HYPOTHESES:
    H3(f"{code} — {enonce}")
    p = doc.add_paragraph(); p.space_after = Pt(2)
    r = p.add_run(f"Résultat : {res}"); r.bold = True; r.font.size = Pt(10)
    r.font.color.rgb = (RGBColor(0x1E,0x7A,0x33) if res == "CONFIRMEE"
                        else RGBColor(0xB0,0x6A,0x00) if res == "PARTIELLEMENT_CONFIRMEE"
                        else RGBColor(0xB0,0x1C,0x1C))
    P("Preuves — " + preuves, 9.5)
    P("Chiffres — " + chiffres, 9.5)
    P("Sources — " + sources, 9, italic=True, color=GREY)
    P("Contradictions et limites — " + contra, 9.5, italic=True)
    P(f"Niveau de confiance : {conf}", 9.5, True)
    P("Implication pour NEXA — " + implic, 9.5, color=BLUE)

# =============================================================== 13. PÉDAGOGIE
doc.add_page_break()
H1("13. Implications pédagogiques")
H2("Bachelor — un socle généraliste réellement technicisé")
P("Le Bachelor doit rester généraliste, parce que c'est la seule façon d'accéder au vivier d'alternance (plus de "
  "1 000 offres) et aux 47,7 % de postes accessibles aux débutants. Mais il ne peut plus être seulement généraliste : "
  "la technicité moyenne de la famille (1,42) le place dans la zone la plus exposée à l'automatisation. Le socle "
  "doit donc être relevé au niveau 2, avec des briques de niveau 3 obligatoires.")
TAB(["Bloc","Contenu","Statut"],
 [["Fondamentaux marketing","Stratégie, positionnement, segmentation, parcours client, business model","Socle"],
  ["Acquisition","SEO, SEA (Google Ads, Microsoft Ads), Paid Social (Meta, LinkedIn, TikTok), bases du programmatique","Socle — obligatoire"],
  ["Contenu et social media","Stratégie éditoriale, création, réseaux sociaux, influence","Socle — mais plus jamais comme spécialisation d'arrivée"],
  ["E-commerce","Gestion de site, marketplaces, e-merchandising, catalogue produits","Socle"],
  ["CRM et automation","CRM (HubSpot en priorité — outil le plus cité), emailing, segmentation, scoring, nurturing","Socle — obligatoire"],
  ["Analytics et tracking","GA4, Google Tag Manager, plan de taggage, Looker Studio, KPI et reporting","Socle — obligatoire, non négociable"],
  ["Data","Excel et Google Sheets avancés, introduction au SQL, lecture de dashboards","Socle — obligatoire"],
  ["Automatisation et IA","Zapier, Make ou n8n ; IA générative encadrée : prompt, contrôle des sorties, gouvernance ; introduction au GEO/AEO","Transversal — obligatoire"],
  ["Gestion de projet et business","Pilotage, budget, ROI, relation client interne, anglais professionnel","Socle"]],
 [3.2,9.0,3.8], 8.5)
P("Compétences devenues trop faibles pour constituer seules une spécialisation : community management, publication "
  "sur les réseaux sociaux, production de contenus génériques, communication digitale institutionnelle. Elles restent "
  "enseignées comme composantes du socle, mais ne peuvent plus donner leur nom à un parcours.", 10, True)

H2("Mastère — deux spécialisations techniques, pas davantage")
P("Les volumes observés ne permettent pas d'ouvrir sept spécialisations. Deux parcours principaux sont soutenus par "
  "le marché, complétés par des modules.")
TAB(["Spécialisation testée","Débouchés observables","Verdict"],
 [["Marketing Performance & Growth","85 offres Acquisition/Performance + 51 SEO/SEA dans l'échantillon ; 1 842 offres Growth Marketing et 318 SEO Manager sur Indeed ; salaires senior 80-90 k€","À OUVRIR — parcours principal"],
  ["CRM, Data & Marketing Automation","36 offres CRM/Lifecycle + 31 Data/Analytics + 17 MarTech ; 2 000+ offres CRM Manager et 759 Marketing Automation sur Indeed ; CRM = 2e compétence du périmètre","À OUVRIR — parcours principal"],
  ["MarTech & AI Marketing","17 offres MarTech et 10 métiers IA émergents seulement ; 0 % d'alternance ; profils de 3 à 7 ans exigés","À NE PAS OUVRIR comme parcours — à intégrer en modules dans les deux parcours principaux"],
  ["E-commerce & Growth","28 offres e-commerce ; secteur à 234 000 emplois (+9 %) ; présent sur tous les territoires","À TRAITER en coloration régionale (Bordeaux, Nantes, Lille, Marseille), pas en parcours national"],
  ["Product Marketing & Revenue","9 offres, 0 % de juniors, 0 % d'alternance, 3 à 10 ans d'expérience exigés","À NE PAS OUVRIR — module optionnel de fin de Mastère"],
  ["Data-Driven Marketing","Recouvre le parcours CRM, Data & Marketing Automation","REDONDANT"],
  ["Growth, CRO & Analytics","Le CRO n'existe quasiment pas comme métier autonome en France (7 mentions) ; c'est une compétence","REDONDANT — CRO à intégrer au parcours Growth"]],
 [4.0,7.5,4.5], 8.5)

H2("L'IA dans le cursus")
BUL("Compétence transversale obligatoire dans les deux parcours, jamais un parcours autonome : la demande employeur observable (2,4 % des offres) ne le justifie pas.")
BUL("Usages suffisamment matures pour être enseignés dès 2026 : génération et retouche de contenus sous contrôle humain, génération créative et déclinaisons publicitaires, automatisation de workflows en no-code, assistance à l'analyse et au reporting, enrichissement CRM, GEO/AEO en complément du SEO.")
BUL("Usages non encore matures côté employeur marketing : orchestration d'agents IA et intégration MCP, aujourd'hui captées par des profils techniques — à enseigner en veille et en découverte, pas en compétence certifiée.")
BUL("Compétences humaines et analytiques dont la valeur augmente : interprétation des données, expérimentation, jugement critique sur les sorties IA, gouvernance et RGPD, compréhension business et pilotage budgétaire.")
BUL("L'IA doit être un socle pédagogique réel et évalué, pas un argument de communication : la prime salariale observée sur l'expertise IA/GEO des profils SEO parisiens (jusqu'à 85-95 k€) montre que le marché sanctionne la compétence réelle.")

H2("Métiers à mettre en avant dans la communication NEXA")
TAB(["À mettre en avant","À dépriorer dans la communication"],
 [["Growth Marketing Manager / Growth Manager","Community Manager"],
  ["Traffic Manager / Consultant Paid Media","Social Media Manager"],
  ["CRM Manager / CRM & Automation Manager","Chargé de communication digitale"],
  ["Consultant SEO / SEA (avec composante GEO)","Content Manager (production)"],
  ["Web Analyst / Marketing Data Analyst","Assistant marketing digital comme finalité"],
  ["E-commerce Manager","Chef de projet digital généraliste"],
  ["Marketing Automation Manager","Influenceur / métiers de l'influence comme débouché principal"]],
 [8.0,8.0], 9)
P("Le raisonnement n'est pas de renoncer aux intitulés populaires par principe, mais de constater que les métiers de "
  "la colonne de droite cumulent la technicité la plus faible, l'exposition la plus forte à l'automatisation et les "
  "salaires les plus bas — alors même qu'ils sont ceux qui attirent le plus de candidatures étudiantes.", 9.5, italic=True)

H2("Référentiel cible et écart avec l'existant")
P("Les programmes actuels de NEXA n'ayant pas été fournis, aucune analyse d'écart contenu par contenu n'a été "
  "réalisée. Le tableau ci-dessous propose un référentiel cible fondé sur la fréquence observée des compétences dans "
  "les offres ; il tient lieu de base pour un audit interne ultérieur.", 10, italic=True)
TAB(["Compétence marché","Fréquence observée","Importance","Niveau requis à la sortie","Urgence"],
 [["Analytics et tracking (GA4, GTM, plan de taggage)","11 offres citées explicitement (borne basse) — présent implicitement dans toutes les offres de niveau 3","Critique","Bachelor : opérationnel · Mastère : avancé (attribution, server-side)","IMMÉDIATE"],
  ["CRM et marketing automation","57 offres (11,5 %) — 2e compétence du périmètre","Critique","Bachelor : opérationnel HubSpot · Mastère : architecture, CDP, cycle de vie","IMMÉDIATE"],
  ["Acquisition payante (Google Ads, Meta, LinkedIn, TikTok)","37 offres SEA + 12 Google Ads + 8 Meta Ads","Critique","Bachelor : opérationnel · Mastère : pilotage budgétaire et attribution","IMMÉDIATE"],
  ["SEO et GEO/AEO","49 offres SEO + 7 GEO/AEO/AI Search","Forte","Bachelor : opérationnel · Mastère : SEO technique et GEO","IMMÉDIATE"],
  ["SQL et exploitation de données","Peu cité explicitement mais exigé dans toutes les offres de Web Analyst et de niveau 4","Forte","Bachelor : introduction · Mastère : autonome (SQL, BigQuery)","ÉLEVÉE"],
  ["Automatisation no-code et IA générative","16 mentions au total sur 495 offres","Émergente mais stratégique","Transversal obligatoire aux deux niveaux","ÉLEVÉE"],
  ["E-commerce et marketplaces","43 offres e-commerce + 8 marketplaces","Forte","Bachelor : opérationnel","MOYENNE"],
  ["CRO et expérimentation","7 offres CRO","Moyenne — compétence, pas métier","Mastère : intégré au parcours Growth","MOYENNE"],
  ["Réseaux sociaux et création de contenu","77 offres (15,6 %) — 1re compétence citée","Forte en volume, faible en valeur","Bachelor uniquement, en socle","À MAINTENIR SANS RENFORCER"]],
 [4.0,4.4,2.4,4.6,2.2], 8)

# =============================================================== 14. SCÉNARIOS
doc.add_page_break()
H1("14. Les quatre scénarios")
for k in ["A_MAINTIEN_GENERALISTE","B_TECHNICISATION","C_HYBRIDE_SPECIALISE","D_REDUCTION_FUSION_FERMETURE"]:
    s = SCENARIOS[k]
    H2(s["titre"])
    TAB(["Critère","Analyse"],
     [["Justification", s["justification"]],
      ["Public cible", s["public_cible"]],
      ["Métiers visés", s["metiers_vises"]],
      ["Volume de débouchés", s["volume_debouches"]],
      ["Accessibilité junior", s["accessibilite_junior"]],
      ["Potentiel d'alternance", s["potentiel_alternance"]],
      ["Compétences", s["competences"]],
      ["Niveau de technicité", s["niveau_technicite"]],
      ["Exposition à l'IA", s["exposition_ia"]],
      ["Résilience à 3-5 ans", s["resilience_3_5_ans"]],
      ["Potentiel national", s["potentiel_national"]],
      ["Potentiel par ville", s["potentiel_par_ville"]],
      ["Avantages", s["avantages"]],
      ["Risques", s["risques"]],
      ["Effort pédagogique", s["effort_pedagogique"]],
      ["Différenciation", s["differenciation"]],
      ["Cohérence avec NEXA", s["coherence_avec_nexa"]],
      ["Niveau de confiance", s["confiance"]],
      ["RECOMMANDATION", s["recommandation"]]],
     [3.4,12.6], 8.5)
    doc.add_paragraph()

# =============================================================== 15. MATRICE
doc.add_page_break()
H1("15. Matrice d'arbitrage")
P("Chaque critère est noté de 0 à 5 pour chacune des quatre orientations, et pondéré : les critères stratégiquement "
  "décisifs pour une école (volume, potentiel à 3-5 ans, accessibilité junior, alternance, résilience à l'IA, "
  "différenciation, pertinence Bachelor, pertinence Mastère) pèsent double. Aucune moyenne mécanique n'a été "
  "appliquée.", 10, italic=True)
rows = []
for nom, poids, notes, just in CRITERES:
    rows.append([nom, f"×{poids}", notes["GENERALISTE"], notes["TECHNIQUE"],
                 notes["HYBRIDE_SPECIALISEE"], notes["REDUCTION_FUSION_FERMETURE"], just])
tot, pmax = scores()
rows.append(["TOTAL PONDÉRÉ", f"/{pmax}", tot["GENERALISTE"], tot["TECHNIQUE"],
             tot["HYBRIDE_SPECIALISEE"], tot["REDUCTION_FUSION_FERMETURE"], ""])
rows.append(["SCORE", "%", f"{round(100*tot['GENERALISTE']/pmax)} %", f"{round(100*tot['TECHNIQUE']/pmax)} %",
             f"{round(100*tot['HYBRIDE_SPECIALISEE']/pmax)} %", f"{round(100*tot['REDUCTION_FUSION_FERMETURE']/pmax)} %", ""])
t = TAB(["Critère","Poids","GÉNÉ-\nRALISTE","TECH-\nNIQUE","HYBRIDE\nSPÉCIALISÉE","RÉDUCTION\nFUSION\nFERMETURE","Documentation du score"],
        rows, [3.0,1.0,1.5,1.5,1.7,1.9,5.4], 7.5)
for i in (len(rows)-1, len(rows)):
    for c in t.rows[i].cells: shade(c, "DEEAF6")
P("")
P(f"Résultat : HYBRIDE SPÉCIALISÉE {round(100*tot['HYBRIDE_SPECIALISEE']/pmax)} % · "
  f"TECHNIQUE {round(100*tot['TECHNIQUE']/pmax)} % · GÉNÉRALISTE {round(100*tot['GENERALISTE']/pmax)} % · "
  f"RÉDUCTION / FUSION / FERMETURE {round(100*tot['REDUCTION_FUSION_FERMETURE']/pmax)} %.", 11, True, color=NAVY)
P("L'écart est net et il ne tient pas à la pondération : le modèle hybride spécialisé l'emporte sur tous les "
  "critères composites parce qu'il est le seul à répondre simultanément aux deux contraintes que le marché impose "
  "— l'alternance et l'accès junior sont dans le généraliste, la valeur et la pérennité sont dans le technique. "
  "L'orientation purement technique arrive deuxième mais échoue sur les deux critères vitaux pour une école : "
  "accessibilité junior (note 1) et alternance (note 1). L'orientation généraliste échoue sur la résilience à l'IA "
  "(note 1), la différenciation (note 1) et la pertinence Mastère (note 1). La fermeture n'est soutenue par aucun "
  "critère de marché.")

# =============================================================== 16. RECOMMANDATION
H1("16. Recommandation finale")
P("ADOPTER LE MODÈLE HYBRIDE SPÉCIALISÉ, avec une inclinaison technique assumée.", 13, True, color=NAVY)
P("Le marché de l'emploi ne justifie ni le maintien d'une filière généraliste inchangée, ni une technicisation "
  "totale, ni la fermeture. Il impose une architecture à deux étages, cohérente avec ce que les données montrent "
  "de façon convergente :")
BUL("Le Bachelor reste généraliste dans son périmètre mais devient technique dans son exigence. Il capte les plus de 1 000 offres d'alternance et les 47,7 % de postes accessibles aux débutants ; il rend obligatoires GA4, GTM, CRM, l'acquisition payante, l'introduction au SQL, l'automatisation no-code et l'IA générative encadrée. Cible de sortie : niveau de technicité 2 confirmé.")
BUL("Le Mastère devient franchement spécialisé et technique, avec deux parcours seulement : Marketing Performance & Growth, et CRM, Data & Marketing Automation. Cible de sortie : niveau de technicité 3, avec des modules de niveau 4 (MarTech, RevOps, CDP, Product Marketing).")
BUL("Le parcours Social Media / Community Management autonome est supprimé : c'est le seul segment de l'étude qui réunit tous les critères de fragilité (technicité 1,03, saturation documentée, réduction des budgets média). Ses contenus sont conservés dans le socle Bachelor.")
BUL("L'IA devient une couche transversale obligatoire et évaluée dans les deux niveaux, et non un parcours affiché — la demande employeur observable (2,4 % des offres) ne soutient pas un parcours dédié.")
BUL("La carte des campus est différenciée : offre complète à Paris, Bachelor plus Mastère Growth à Lyon, Bachelor seul avec coloration locale à Lille, Bordeaux, Marseille et Nantes, Mastères spécialisés accessibles en distanciel national.")
BUL("Des passerelles et des modules communs sont ouverts avec la filière IA & Data (data marketing, SQL, dashboards) et la filière Développement Web (tracking, intégrations, API), qui constituent l'avantage différenciant de NEXA face aux écoles de commerce.")

# =============================================================== SOURCES
doc.add_page_break()
H1("Études, articles et sources pour approfondir")
P("Toutes les URL ci-dessous ont été consultées le 7 septembre 2026. Les niveaux de fiabilité sont ceux retenus "
  "pour l'analyse : FORTE pour les organismes producteurs de données primaires (Apec, France Travail, Dares, Fevad, "
  "SRI, Alliance Digitale/EY), MOYENNE pour les reprises documentées et les études de cabinets, FAIBLE pour les "
  "compilations sectorielles secondaires.", 9.5, italic=True, color=GREY)
ordre = {"FORTE": 0, "MOYENNE": 1, "FAIBLE": 2}
for e in sorted(ETUDES, key=lambda x: (ordre.get(x["NIVEAU_DE_FIABILITE"], 3), x["ORGANISME"])):
    p = doc.add_paragraph(); p.space_after = Pt(1)
    r = p.add_run(e["TITRE"]); r.bold = True; r.font.size = Pt(9.5)
    P(f"{e['ORGANISME']} · publié : {e['DATE_PUBLICATION']} · consulté : {e['DATE_CONSULTATION']} · "
      f"périmètre : {e['PERIMETRE']} · méthode : {e['METHODE']} · échantillon : {e['TAILLE_ECHANTILLON']} · "
      f"page : {e['PAGE_DU_RAPPORT']} · fiabilité : {e['NIVEAU_DE_FIABILITE']}", 8.5, color=GREY, after=1)
    P("Résultat utilisé — " + e["RESULTAT_UTILISE"], 8.5, after=1)
    p = doc.add_paragraph(); p.space_after = Pt(7)
    link(p, e["URL"])

P("Sources d'offres et de volumes", 11, True)
P("Les 507 offres et les 90 stocks d'offres datés proviennent des plateformes suivantes, dont les URL complètes "
  "figurent dans l'onglet OFFRES_DETAILLEES et dans le bloc STOCKS D'OFFRES DATÉS du fichier Excel joint :", 9.5)
for nom, url in [("Indeed France", "https://fr.indeed.com/"),
                 ("Welcome to the Jungle", "https://www.welcometothejungle.com/fr"),
                 ("HelloWork", "https://www.hellowork.com/fr-fr/"),
                 ("France Travail", "https://candidat.francetravail.fr/offres/emploi")]:
    p = doc.add_paragraph(style="List Bullet"); p.space_after = Pt(1)
    p.add_run(nom + " — ").font.size = Pt(9)
    link(p, url)

# =============================================================== ENCADRÉ FINAL
doc.add_page_break()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.space_after = Pt(10)
r = p.add_run("DÉCISION RECOMMANDÉE POUR NEXA"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NAVY

def box(label, contenu, gras=False, fill="FFFFFF"):
    t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"
    t.rows[0].cells[0].width = Cm(5.0); t.rows[0].cells[1].width = Cm(11.5)
    c0 = t.rows[0].cells[0]; c0.text = ""
    r0 = c0.paragraphs[0].add_run(label); r0.bold = True; r0.font.size = Pt(9)
    r0.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); shade(c0, "1F3864")
    c1 = t.rows[0].cells[1]; c1.text = ""; shade(c1, fill)
    r1 = c1.paragraphs[0].add_run(contenu); r1.font.size = Pt(9.5); r1.bold = gras
    return t

box("ORIENTATION RECOMMANDÉE", "HYBRIDE_SPECIALISEE", True, "DEEAF6")
box("DÉCISION",
 "Conserver la filière Marketing Digital comme filière autonome, mais la restructurer en deux étages : un Bachelor "
 "généraliste réellement technicisé (cible de technicité 2, socle obligatoire GA4, GTM, CRM, acquisition payante, "
 "introduction au SQL, automatisation no-code, IA générative encadrée), et un Mastère franchement spécialisé et "
 "technique avec deux parcours seulement — Marketing Performance & Growth, et CRM, Data & Marketing Automation. "
 "Supprimer le parcours Social Media / Community Management autonome. Traiter l'IA comme couche transversale "
 "obligatoire et évaluée, jamais comme parcours affiché.", True, "DEEAF6")
box("JUSTIFICATION EN 5 POINTS",
 "1. Le marché croît et ne justifie aucune fermeture : 310 000 emplois de filière, croissance cinq fois supérieure "
 "au PIB, publicité digitale +11 % en 2025 et +12 % au S1 2026, e-commerce +9 % d'emplois, +4 % de recrutements de "
 "cadres attendus en 2026.\n"
 "2. Le marché est majoritairement spécialisé : 51,9 % des offres du périmètre relèvent des familles spécialisées "
 "contre 22,4 % pour le cœur généraliste.\n"
 "3. Mais les débouchés d'entrée sont dans le généraliste : 40,5 % d'alternance et 47,7 % de débutants et juniors, "
 "contre 0 % d'alternance en MarTech, en Product Marketing et sur les métiers IA émergents, et 3,2 % en data "
 "marketing. Un Bachelor purement technique n'aurait ni alternance ni placement.\n"
 "4. La valeur, les salaires et la résilience à l'automatisation sont dans le technique : technicité de 2,09 à 2,82 "
 "sur les familles spécialisées contre 1,42 sur le généraliste, et un écart salarial d'un facteur trois entre "
 "l'exécution généraliste et l'expertise senior.\n"
 "5. Le marketeur augmenté par l'IA n'est pas encore un marché : 2,4 % des offres mentionnent explicitement l'IA, "
 "un chiffre qui converge avec la mesure Apec (2 % des offres commercial-marketing), et les postes d'orchestration "
 "d'agents IA sont aujourd'hui captés par des profils techniques.")
box("MÉTIERS À PRIORISER",
 "Growth Marketing Manager · Traffic Manager et Consultant Paid Media · CRM Manager et CRM & Automation Manager · "
 "Marketing Automation Manager · Consultant SEO/SEA avec composante GEO · Web Analyst et Marketing Data Analyst · "
 "E-commerce Manager · Responsable Acquisition")
box("MÉTIERS À DÉPRIORISER",
 "Community Manager · Social Media Manager · Chargé de communication digitale · Content Manager centré sur la "
 "production · Assistant marketing digital présenté comme finalité · Chef de projet digital généraliste · "
 "Métiers de l'influence comme débouché principal")
box("COMPÉTENCES À RENFORCER",
 "GA4, Google Tag Manager, plan de taggage et attribution · CRM (HubSpot en priorité) et marketing automation · "
 "acquisition payante et pilotage budgétaire · SQL et exploitation de données, BigQuery, dashboards · CRO et "
 "expérimentation · automatisation no-code (Zapier, Make, n8n) · IA générative encadrée, contrôle des sorties et "
 "gouvernance · GEO/AEO · RGPD et qualité de la donnée · compréhension business et ROI")
box("COMPÉTENCES À RENDRE SECONDAIRES",
 "Community management opérationnel · publication et animation des réseaux sociaux · production de contenus "
 "génériques · communication institutionnelle et relations presse · création visuelle non marketing · "
 "coordination sans pilotage de la performance")
box("POSITIONNEMENT RECOMMANDÉ DU BACHELOR",
 "Bachelor Marketing Digital & Acquisition — socle généraliste technicisé, cible de sortie au niveau de technicité 2 "
 "confirmé, avec briques de niveau 3 obligatoires (GA4, GTM, CRM, introduction SQL). Déployable sur les six campus. "
 "Objectif : capter le vivier d'alternance et garantir le placement.")
box("POSITIONNEMENT RECOMMANDÉ DU MASTÈRE",
 "Deux parcours seulement, cible de sortie au niveau de technicité 3 avec modules de niveau 4 :\n"
 "· Mastère Marketing Performance & Growth (acquisition, paid, SEO/GEO, CRO, analytics, pilotage budgétaire)\n"
 "· Mastère CRM, Data & Marketing Automation (CRM, lifecycle, automation, CDP, SQL, dashboards, MarTech, RevOps)\n"
 "Modules transversaux : IA appliquée au marketing, Product Marketing, e-commerce, RGPD et gouvernance de la donnée.")
box("IMPACT PAR CAMPUS",
 "Paris / Île-de-France (2 000+ offres) : offre complète, les deux Mastères, modules MarTech et Product Marketing.\n"
 "Lyon (258 à 317 offres) : Bachelor + Mastère Marketing Performance & Growth.\n"
 "Lille (113 offres) : Bachelor avec coloration CRM et e-commerce ; Mastère en distanciel.\n"
 "Marseille-Aix (92 offres) : Bachelor avec coloration e-commerce et social ads ; Mastère en distanciel.\n"
 "Bordeaux (90 offres) : Bachelor avec coloration e-commerce et acquisition ; Mastère en distanciel.\n"
 "Nantes (82 offres, marché le plus étroit) : Bachelor seul, coloration e-commerce et SEO/SEA ; ne pas ouvrir de "
 "Mastère spécialisé en présentiel.\n"
 "Distanciel national : levier décisif pour rentabiliser les Mastères spécialisés à faible volume local.")
box("RISQUE SI NEXA NE CHANGE RIEN",
 "Une filière indifférenciable de dizaines d'offres concurrentes, formant majoritairement à des compétences de "
 "niveau de technicité 1 — celles que 75 à 88 % des marketeurs déclarent déjà automatiser avec l'IA générative — "
 "pour des salaires de sortie de 20 000 à 30 000 €, sur un segment social media saturé où 52 % des annonceurs "
 "réduisent leur budget média. Le risque n'est pas l'effondrement immédiat du recrutement étudiant : il est "
 "l'érosion progressive du taux de placement en sortie et de la valeur perçue du diplôme, dans un marché qui, lui, "
 "continue de croître et de se techniciser, au profit d'écoles mieux positionnées.", False, "FCE4EC")
box("NIVEAU DE CONFIANCE",
 "ÉLEVÉ sur la structure du marché, la hiérarchie des familles de métiers, la répartition de l'alternance et de "
 "l'accès junior, la géographie et le diagnostic sur l'IA (convergence de l'échantillon et de la mesure Apec).\n"
 "MOYEN sur les salaires par métier (13 offres seulement affichent une rémunération exploitable) et sur la "
 "différenciation concurrentielle (aucune analyse de l'offre de formation n'a été réalisée).\n"
 "FAIBLE sur les évolutions année contre année par métier : les séries historiques comparables n'ont pas pu être "
 "reconstituées et cette limite est signalée partout où elle s'applique.", False, "FFF2CC")
box("DONNÉES COMPLÉMENTAIRES NÉCESSAIRES",
 "Cette recommandation porte UNIQUEMENT sur la pertinence de la filière au regard du marché de l'emploi. Une "
 "décision économique définitive exige les indicateurs internes NEXA qui n'ont pas été fournis : candidatures et "
 "inscriptions par parcours, taux de remplissage, nombre et qualité des contrats d'alternance effectivement signés "
 "par campus, taux de placement à 6 et 12 mois, salaires de sortie constatés, taux d'abandon, marge et coût "
 "pédagogique par parcours. Deux compléments d'étude externes seraient également utiles : une analyse "
 "concurrentielle des programmes marketing digital français (Bachelor et Mastère), et un accès à l'API France "
 "Travail ou aux données Apec par métier pour reconstituer de véritables séries annuelles.", False, "FFF2CC")

path = os.path.join(ROOT, "NEXA_Synthese_Marche_Emploi_Marketing_Digital_2026.docx")
doc.save(path)
print("Écrit :", path)
